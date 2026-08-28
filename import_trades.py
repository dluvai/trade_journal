"""Parsing/mapping helpers for importing trades from an arbitrary CSV or
XLSX file -- unlike the app's own Export CSV format, a file from "anywhere"
can't be auto-parsed reliably, so the flow is: parse a sample, let the user
confirm/adjust a column mapping, then parse+insert the full file using that
mapping. No new dependency for XLSX (openpyxl is already used elsewhere in
this app for the original journal import) or for date parsing (a fixed list
of formats tried in order, stdlib only).
"""
import csv
import io
from datetime import date, datetime

import openpyxl

# Single source of truth for the mapping UI; "result" is deliberately absent since it's always server-derived from rr.
TARGET_FIELDS = [
    ("date", "Date"), ("session", "Session"), ("pair", "Pair"), ("direction", "Direction"),
    ("risk", "Risk %"), ("rr", "RR"), ("pnl", "PnL %"), ("notes", "Notes"),
    ("chart_daily", "Daily Chart Link"), ("chart_4h", "4H Chart Link"), ("chart_30m", "30M Chart Link"),
    ("ignore", "Ignore"),
]
TARGET_FIELD_KEYS = {k for k, _ in TARGET_FIELDS}

SUGGESTION_KEYWORDS = {
    "date": ["date", "time"],
    "pair": ["symbol", "instrument", "pair", "ticker", "market"],
    "direction": ["side", "direction", "type", "action"],
    "session": ["session"],
    "risk": ["risk"],
    "rr": ["rr", "r multiple", "r-multiple", "reward"],
    "pnl": ["pnl", "p/l", "profit", "return", "gain"],
    "notes": ["notes", "comment", "remarks", "reason"],
    "chart_daily": ["daily chart", "chart daily"],
    "chart_4h": ["4h chart", "chart 4h"],
    "chart_30m": ["30m chart", "chart 30m"],
}

DATE_FORMATS = ["%Y-%m-%d", "%m/%d/%Y", "%d/%m/%Y", "%Y/%m/%d", "%m-%d-%Y", "%d-%m-%Y"]

DIRECTION_SYNONYMS = {
    "buy": "Long", "long": "Long", "bullish": "Long", "b": "Long",
    "sell": "Short", "short": "Short", "bearish": "Short", "s": "Short",
}


def suggest_mapping(header_row):
    mapping = {}
    for i, col in enumerate(header_row):
        col_lower = str(col).strip().lower()
        best = "ignore"
        for field, keywords in SUGGESTION_KEYWORDS.items():
            if any(kw in col_lower for kw in keywords):
                best = field
                break
        mapping[str(i)] = best
    return mapping


def looks_like_header(first_row, second_row):
    # Heuristic-only header guess (non-numeric row 1, numeric/date row 2) -- the user can always override it.
    if not second_row:
        return True

    def is_numeric_or_date(v):
        s = str(v).strip()
        if not s:
            return False
        try:
            float(s)
            return True
        except ValueError:
            pass
        return parse_date_flexible(s) is not None

    first_numeric = sum(1 for v in first_row if is_numeric_or_date(v))
    second_numeric = sum(1 for v in second_row if is_numeric_or_date(v))
    return first_numeric < second_numeric


def parse_csv_sample(file_bytes, n=10):
    text = file_bytes.decode("utf-8-sig")
    reader = list(csv.reader(io.StringIO(text)))
    rows = reader[: n + 1]
    return _shape_sample(rows)


def parse_xlsx_sample(file_bytes, n=10):
    wb = openpyxl.load_workbook(io.BytesIO(file_bytes), read_only=True, data_only=True)
    ws = wb.worksheets[0]
    rows = []
    for i, row in enumerate(ws.iter_rows(values_only=True)):
        if i > n:
            break
        rows.append(["" if v is None else v for v in row])
    return _shape_sample(rows)


def _shape_sample(rows):
    if not rows:
        return {"columns": [], "sample_rows": [], "has_header_guess": True, "suggested_mapping": {}}
    header_guess = looks_like_header(rows[0], rows[1] if len(rows) > 1 else None)
    columns = [str(c).strip() or f"Column {i+1}" for i, c in enumerate(rows[0])]
    return {
        "columns": columns,
        "sample_rows": rows[:5],
        "has_header_guess": header_guess,
        "suggested_mapping": suggest_mapping(rows[0]),
    }


def parse_full_csv(file_bytes):
    text = file_bytes.decode("utf-8-sig")
    return list(csv.reader(io.StringIO(text)))


def parse_full_xlsx(file_bytes):
    wb = openpyxl.load_workbook(io.BytesIO(file_bytes), read_only=True, data_only=True)
    ws = wb.worksheets[0]
    return [["" if v is None else v for v in row] for row in ws.iter_rows(values_only=True)]


def parse_date_flexible(raw):
    # Checked before string-format parsing since openpyxl returns real date cells as datetime objects, not strings.
    if isinstance(raw, datetime):
        return raw.date().isoformat()
    if isinstance(raw, date):
        return raw.isoformat()
    raw = str(raw).strip()
    for fmt in DATE_FORMATS:
        try:
            return datetime.strptime(raw, fmt).strftime("%Y-%m-%d")
        except ValueError:
            continue
    return None


def normalize_direction(raw):
    return DIRECTION_SYNONYMS.get(str(raw).strip().lower(), str(raw).strip())


def row_to_trade_fields(row, mapping, risk_pnl_is_percent):
    # "result" can never be mapped since TARGET_FIELDS has no such entry.
    raw = {}
    for idx_str, field in mapping.items():
        if field == "ignore" or field not in TARGET_FIELD_KEYS:
            continue
        idx = int(idx_str)
        if idx < len(row):
            raw[field] = row[idx]

    fields = {}
    if "date" in raw:
        parsed = parse_date_flexible(raw["date"])
        if parsed is None:
            raise ValueError(f"Unrecognized date: {raw['date']!r}")
        fields["date"] = parsed
    if "session" in raw:
        fields["session"] = str(raw["session"]).strip()
    if "pair" in raw:
        fields["pair"] = str(raw["pair"]).strip().upper()
    if "direction" in raw:
        fields["direction"] = normalize_direction(raw["direction"])
    if "notes" in raw:
        fields["notes"] = str(raw["notes"]).strip()
    for chart_key in ("chart_daily", "chart_4h", "chart_30m"):
        if chart_key in raw:
            fields[chart_key] = str(raw[chart_key]).strip()

    for numeric_key in ("risk", "rr", "pnl"):
        if numeric_key not in raw or str(raw[numeric_key]).strip() == "":
            continue
        val = float(str(raw[numeric_key]).strip().rstrip("%"))
        # rr is a plain multiple, not percent-scaled like risk/pnl.
        if numeric_key in ("risk", "pnl") and risk_pnl_is_percent:
            val = val / 100
        fields[numeric_key] = val

    return fields
